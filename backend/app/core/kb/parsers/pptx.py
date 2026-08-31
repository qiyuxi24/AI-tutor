"""
PPT 演示文稿解析器

支持: .pptx（Office Open XML 新格式）
方案：python-pptx 提取每张幻灯片的文本：
- 标题 + 各文本框段落
- 表格单元格
- 备注页文本（可选，作为补充上下文）
注：老式二进制 .ppt 需专用方案，见 legacy.py（可选）。
"""

from __future__ import annotations

import io

from app.core.kb.parsers.base import BaseParser, ParseResult


class PptxParser(BaseParser):
    """PPT .pptx 解析器（python-pptx）"""

    name = "pptx"
    extensions = {"pptx"}

    def parse(self, filename: str, content: bytes) -> ParseResult:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(content))
        slides_text: list[str] = []

        for i, slide in enumerate(prs.slides, start=1):
            lines: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        txt = "".join(run.text for run in para.runs).strip()
                        if txt:
                            lines.append(txt)
                elif shape.has_table:
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        lines.append(" | ".join(cells))
            # 备注页（若存在）
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    lines.append(f"[备注] {notes}")

            if lines:
                slides_text.append(f"【第 {i} 页】\n" + "\n".join(lines))

        return self._ok("\n\n".join(slides_text), meta={"slides": len(prs.slides)})
