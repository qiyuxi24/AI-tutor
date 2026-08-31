"""
Word 文档解析器

支持: .docx（Office Open XML 新格式）
方案：python-docx 提取段落 + 表格文本。
注：老式二进制 .doc 需专用方案，见 legacy.py（可选）。
"""

from __future__ import annotations

import io

from app.core.kb.parsers.base import BaseParser, ParseResult


class DocxParser(BaseParser):
    """Word .docx 解析器（python-docx）"""

    name = "docx"
    extensions = {"docx"}

    def parse(self, filename: str, content: bytes) -> ParseResult:
        from docx import Document
        doc = Document(io.BytesIO(content))
        parts: list[str] = []
        # 段落文本
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text)
        # 表格文本
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                parts.append(" | ".join(cells))
        return self._ok("\n".join(parts))
